import os
import json
import uuid
import aiohttp
import asyncio
import requests
import base64
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
import whisper
import numpy as np
import soundfile as sf
import io
import subprocess
import tempfile
import PyPDF2
import docx
import openpyxl
import re
from pathlib import Path
from threading import Lock
import asyncio
import aiohttp
from sentence_transformers import SentenceTransformer, util
import numpy as np

OLLAMA_TIMEOUT = 60  # Increased timeout to 60 seconds

# ========== CONFIG ==========
OLLAMA_URL = "http://localhost:11434/api/chat"
MODEL_NAME = "mistral"
CONVERSATION_FILE = "conversations.json"

# Load Whisper model (small is fast, runs locally)
whisper_model = whisper.load_model("small")
import threading
import time



async def fetch_ai_reply(session, messages, retries=3):
    for attempt in range(retries):
        try:
            print(f"[DEBUG] Attempting AI reply call {attempt + 1}/{retries}")
            async with session.post(OLLAMA_URL, json={
                "model": MODEL_NAME,
                "messages": messages,
                "stream": False
            }, timeout=OLLAMA_TIMEOUT) as response:
                if response.ok:
                    data = await response.json()
                    content = data.get("message", {}).get("content", "❌ Error from LLM backend.")
                    print(f"[DEBUG] AI reply received: {content[:100]}...")
                    return content
                else:
                    print(f"[ERROR] LLM backend returned status {response.status} on attempt {attempt+1}")
                    await asyncio.sleep(1)
        except asyncio.TimeoutError:
            print(f"[ERROR] LLM backend call timed out on attempt {attempt+1}")
            await asyncio.sleep(2)
        except Exception as e:
            print(f"[ERROR] LLM backend call error on attempt {attempt+1}: {e}")
            await asyncio.sleep(1)
    return "⚠️ Unable to reach backend after retries. Please try again."


async def fetch_real_mcp_reply(session, user_question, max_retries=5):
    mcp_prompt = [
        {"role": "system", "content": """You are a strict MCP assistant. Reply ONLY with JSON with keys: search_needed (bool), search_query (string|null), assistant_reply (string|null). No text outside JSON.

RULES:
- Set search_needed to FALSE for: greetings, casual conversation, general questions, personal questions, jokes, etc.
- Set search_needed to TRUE only for: specific technical questions, questions about documents/files, questions requiring factual information from stored data
- For search_needed=false, provide a helpful assistant_reply and set search_query to null
- For search_needed=true, set search_query to the specific terms to search for and assistant_reply to null

Examples:
- "Hello" → {"search_needed": false, "search_query": null, "assistant_reply": "Hello! How can I help you today?"}
- "How are you?" → {"search_needed": false, "search_query": null, "assistant_reply": "I'm doing well, thank you for asking! How can I assist you?"}
- "What's in the VESG-2 document?" → {"search_needed": true, "search_query": "VESG-2 document content", "assistant_reply": null}"""},
        {"role": "user", "content": user_question}
    ]

    for attempt in range(max_retries):
        try:
            print(f"[DEBUG] Attempting MCP call {attempt + 1}/{max_retries}")
            async with session.post(OLLAMA_URL, json={"model": MODEL_NAME, "messages": mcp_prompt, "stream": False}, timeout=OLLAMA_TIMEOUT) as resp:
                if not resp.ok:
                    print(f"[ERROR] Ollama returned status {resp.status} on attempt {attempt + 1}")
                    await asyncio.sleep(1)  # Wait before retry
                    continue
                
                data = await resp.json()
                content = data.get("message", {}).get("content", "")
                print(f"[DEBUG] Received content: {content[:100]}...")
                
                try:
                    parsed = json.loads(content)
                    # Validate all fields strictly
                    if (
                        isinstance(parsed, dict) and
                        isinstance(parsed.get("search_needed"), bool) and
                        (isinstance(parsed.get("search_query"), (str, type(None)))) and
                        (isinstance(parsed.get("assistant_reply"), (str, type(None))))
                    ):
                        print(f"[DEBUG] Valid MCP response received")
                        return parsed
                    else:
                        print(f"[DEBUG] Invalid MCP response structure, retrying...")
                        # Prepare a strict correction prompt to send to model
                        correction_msg = {
                            "role": "system",
                            "content": "Your last reply was not a valid MCP JSON. Reply ONLY with valid JSON with correct keys and types."
                        }
                        mcp_prompt.append({"role": "assistant", "content": content})
                        mcp_prompt.append(correction_msg)
                except json.JSONDecodeError as e:
                    print(f"[DEBUG] JSON decode error: {e}, retrying...")
                    # Same as above: send correction prompt
                    correction_msg = {
                        "role": "system",
                        "content": "Your last reply was invalid JSON. Reply ONLY with valid MCP JSON."
                    }
                    mcp_prompt.append({"role": "assistant", "content": content})
                    mcp_prompt.append(correction_msg)
                    
        except asyncio.TimeoutError:
            print(f"[ERROR] Ollama request timed out on attempt {attempt + 1}")
            await asyncio.sleep(2)  # Wait longer after timeout
        except Exception as e:
            print(f"[ERROR] Unexpected error on attempt {attempt + 1}: {e}")
            await asyncio.sleep(1)

    # After max retries, fail gracefully
    print(f"[ERROR] MCP protocol failed after {max_retries} retries")
    return {
        "search_needed": False,
        "search_query": None,
        "assistant_reply": "⚠️ MCP protocol failed after retries. Please try again."
    }





def background_file_cache_refresher(interval_seconds=3600):
    global file_cache
    while True:
        print("[Background] Refreshing file cache...")
        new_cache = {}
        p = Path(FILE_FOLDER)
        for file_path in p.glob("*"):
            ext = file_path.suffix.lower()
            try:
                if ext == ".pdf":
                    text = extract_text_from_pdf(file_path)
                elif ext == ".docx":
                    text = extract_text_from_docx(file_path)
                elif ext == ".xlsx":
                    text = extract_text_from_xlsx(file_path)
                elif ext == ".txt":
                    text = extract_text_from_txt(file_path)
                else:
                    continue
                new_cache[str(file_path)] = text
            except Exception as e:
                print(f"[Background] Error loading {file_path}: {e}")
        with file_cache_lock:
            file_cache = new_cache
        refresh_file_embeddings()  # Update embeddings
        print(f"[Background] File cache refreshed with {len(file_cache)} files.")
        time.sleep(interval_seconds)


# Load model for semantic search
embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

# Cache for precomputed embeddings
file_embeddings = {}  # {path: {"chunks": [], "embeddings": []}}

def chunk_text(text, max_length=100, overlap=20):
    words = text.split()
    chunks = []
    for i in range(0, len(words), max_length - overlap):
        chunk = " ".join(words[i:i + max_length])
        chunks.append(chunk)
    return chunks

def refresh_file_embeddings():
    global file_embeddings
    texts = load_all_files()
    file_embeddings = {}
    for path, content in texts.items():
        chunks = chunk_text(content)
        embeddings = embedding_model.encode(chunks, convert_to_tensor=True)
        file_embeddings[path] = {"chunks": chunks, "embeddings": embeddings}

def semantic_search_files(query, top_k=3):
    query_embedding = embedding_model.encode(query, convert_to_tensor=True)
    results = []

    for path, data in file_embeddings.items():
        cos_scores = util.cos_sim(query_embedding, data["embeddings"]).cpu().numpy()[0]
        top_indices = np.argsort(cos_scores)[-top_k:][::-1]

        for idx in top_indices:
            snippet = data["chunks"][idx]
            results.append(f"...{snippet.strip()}...")

    return "\n".join(results) if results else None



# ========== FLASK SETUP ==========
app = Flask(__name__)
CORS(app)

# ========== UTILITY FUNCTIONS ==========
def load_conversations():
    if os.path.exists(CONVERSATION_FILE):
        try:
            with open(CONVERSATION_FILE, "r") as f:
                data = json.load(f)
                print(f"[DEBUG] Loaded {len(data)} conversations")
                return data
        except Exception as e:
            print(f"Error loading conversations: {e}")
            return {}
    print(f"[DEBUG] No conversation file found at {CONVERSATION_FILE}")
    return {}

def save_conversations(data):
    print(f"Saving conversations to: {os.path.abspath(CONVERSATION_FILE)}")
    print(f"Current working directory: {os.getcwd()}")
    try:
        with open(CONVERSATION_FILE, "w") as f:
            json.dump(data, f, indent=2)
            f.flush()
            os.fsync(f.fileno())
        print(f"Data length: {len(json.dumps(data))}")
        print(f"Conversations saved successfully")
    except Exception as e:
        print(f"Error saving conversations: {e}")



# ========== FILE SEARCH HELPERS ==========

FILE_FOLDER = r"C:\Users\user\Desktop\LPEE BOT\backend\files"
file_cache = {}
file_cache_lock = Lock()

def extract_text_from_pdf(path):
    text = []
    try:
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                raw = page.extract_text()
                if raw:
                    # Normalize spacing and fix broken lines
                    raw = re.sub(r'(?<!\.)\s+$', '', raw.strip())  # remove trailing space
                    raw = re.sub(r'(\w+)-\n(\w+)', r'\1\2', raw)  # fix hyphenated words split across lines
                    raw = re.sub(r'([a-zA-Z])- ([a-zA-Z])', r'\1\2', raw)  # fix word breaks
                    text.append(raw)
    except Exception as e:
        print(f"PDF extraction error for {path}: {e}")
    return "\n".join(text)

def extract_text_from_docx(path):
    text = []
    try:
        doc = docx.Document(path)
        for para in doc.paragraphs:
            text.append(para.text)
    except Exception as e:
        print(f"DOCX extraction error for {path}: {e}")
    return "\n".join(text)

def extract_text_from_xlsx(path):
    text = []
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = " ".join([str(cell) if cell is not None else "" for cell in row])
                text.append(line)
    except Exception as e:
        print(f"XLSX extraction error for {path}: {e}")
    return "\n".join(text)

def extract_text_from_txt(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except Exception as e:
        print(f"TXT extraction error for {path}: {e}")
        return ""

def load_all_files():
    global file_cache
    with file_cache_lock:
        return file_cache


def search_files_for_answer(query):
    texts = load_all_files()
    results = []
    pattern = re.compile(re.escape(query), re.IGNORECASE)
    for path, content in texts.items():
        if pattern.search(content):
            snippet_start = content.lower().find(query.lower())
            snippet_end = snippet_start + 300 if snippet_start >= 0 else 300
            snippet = content[snippet_start:snippet_end].replace("\n", " ").strip()
            results.append(f"From {path}: ...{snippet}...")
    return "\n\n".join(results) if results else None





# ========== API ROUTES ==========

@app.route('/api/stt', methods=['POST'])
def stt():
    try:
        audio_file = request.files['audio']
        if not audio_file:
            return jsonify({ 'error': 'No audio file uploaded' }), 400

        # Save uploaded audio to a temporary file
        with tempfile.NamedTemporaryFile(suffix=".webm", delete=False) as temp_input:
            temp_input.write(audio_file.read())
            temp_input.flush()
            input_path = temp_input.name

        # Convert WebM to WAV using ffmpeg subprocess
        with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as temp_wav:
            output_path = temp_wav.name

        subprocess.run([
            "ffmpeg", "-y", "-i", input_path, "-ar", "16000", "-ac", "1", output_path
        ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

        # Load WAV and transcribe
        audio_data, sr = sf.read(output_path)
        result = whisper_model.transcribe(output_path, fp16=False)
        text = result['text'].strip()

        return jsonify({ 'text': text })

    except Exception as e:
        print("STT error:", e)
        return jsonify({ 'error': 'Error converting speech to text' }), 500

@app.route('/api/chat', methods=['POST'])
async def chat():
    try:
        data = request.get_json()
        messages = data.get("messages", [])
        conversation_id = data.get("conversationId")

        if not conversation_id or not messages:
            return jsonify({"error": "Invalid conversation data"}), 400

        user_question = messages[-1]['content']
        print(f"[DEBUG] User question: {user_question}")

        # Quick check for simple greetings/conversation that don't need file search
        simple_greetings = ['hello', 'hi', 'hey', 'how are you', 'good morning', 'good afternoon', 'good evening', 'thanks', 'thank you']
        is_simple_greeting = any(greeting in user_question.lower() for greeting in simple_greetings)
        
        if is_simple_greeting:
            print(f"[DEBUG] Detected simple greeting, skipping MCP and file search")
            # Generate a simple response without MCP
            simple_prompt = [
                {"role": "system", "content": "You are a friendly and helpful assistant. Respond naturally to greetings and casual conversation."},
                {"role": "user", "content": user_question}
            ]
            
            async with aiohttp.ClientSession() as session:
                final_reply = await fetch_ai_reply(session, simple_prompt)
        else:
            # Use MCP for more complex queries
            async with aiohttp.ClientSession() as session:
                try:
                    mcp_response = await fetch_real_mcp_reply(session, user_question)
                    print(f"[DEBUG] MCP response: {mcp_response}")

                    final_reply = mcp_response.get("assistant_reply") or ""

                    if mcp_response.get("search_needed"):
                        query = mcp_response["search_query"]
                        print(f"[DEBUG] MCP decided to search files with query: {query}")

                        semantic_result = semantic_search_files(query)
                        if semantic_result:
                            file_search_result = semantic_result
                        else:
                            file_search_result = search_files_for_answer_loose(query)

                        if file_search_result:
                            print(f"[DEBUG] File search found relevant content, enriching reply...")

                            enhanced_prompt = [
                                {
                                    "role": "system",
                                    "content": (
                                        "You are a helpful assistant answering based ONLY on the following context. "
                                        "Do NOT mention documents, excerpts, or context in your answer. Just answer as if you knew it directly."
                                        f"{file_search_result}"
                                    )
                                },
                                {
                                    "role": "system",
                                    "content": f"Context:\n{file_search_result}"
                                },
                                {
                                    "role": "user",
                                    "content": user_question
                                }
                            ]

                            final_reply = await fetch_ai_reply(session, enhanced_prompt)
                            print(f"[DEBUG] Final enhanced reply with docs: {final_reply}")
                except Exception as e:
                    print(f"[ERROR] Error in MCP processing: {e}")
                    # Fallback to direct AI response
                    fallback_prompt = [
                        {"role": "system", "content": "You are a helpful assistant."},
                        {"role": "user", "content": user_question}
                    ]
                    final_reply = await fetch_ai_reply(session, fallback_prompt)

        convos = load_conversations()
        if conversation_id not in convos:
            return jsonify({"error": "Conversation not found"}), 404

        convos[conversation_id]["messages"].append({"role": "user", "content": user_question})
        convos[conversation_id]["messages"].append({"role": "assistant", "content": final_reply})
        save_conversations(convos)

        return jsonify({"content": final_reply})
                
    except Exception as e:
        print(f"[ERROR] Error in chat endpoint: {e}")
        return jsonify({"error": "Internal server error"}), 500



@app.route('/api/health', methods=['GET'])
def health_check():
    """Health check endpoint to test Ollama connectivity"""
    try:
        response = requests.get("http://localhost:11434/api/tags", timeout=10)
        if response.ok:
            return jsonify({"status": "healthy", "ollama": "connected"})
        else:
            return jsonify({"status": "unhealthy", "ollama": "error", "status_code": response.status_code}), 500
    except requests.exceptions.Timeout:
        return jsonify({"status": "unhealthy", "ollama": "timeout"}), 500
    except requests.exceptions.ConnectionError:
        return jsonify({"status": "unhealthy", "ollama": "connection_error"}), 500
    except Exception as e:
        return jsonify({"status": "unhealthy", "ollama": "error", "message": str(e)}), 500

@app.route('/api/refresh-embeddings', methods=['POST'])
def refresh_embeddings():
    print("[Manual Refresh] Refreshing embeddings...")
    refresh_file_embeddings()
    return jsonify({"status": "Embeddings refreshed"})




# Helper: Loose matching file search function
def search_files_for_answer_loose(query, context_chars=300):
    texts = load_all_files()
    results = []
    query_lower = query.lower()
    words = set(w for w in re.findall(r'\w+', query_lower) if len(w) > 2)
    if not words:
        return None

    for path, content in texts.items():
        content_lower = content.lower()
        # Find first occurrence of any keyword
        idx = -1
        for word in words:
            i = content_lower.find(word)
            if i != -1 and (idx == -1 or i < idx):
                idx = i
        if idx == -1:
            continue
        start = max(0, idx - context_chars)
        end = min(len(content), idx + context_chars)
        snippet = content[start:end].replace("\n", " ").strip()
        results.append(f"...{snippet}...")
    return "\n\n".join(results) if results else None




@app.route('/api/test-files', methods=['GET'])
def test_files():
    files = load_all_files()
    return jsonify({
        "file_count": len(files),
        "files": list(files.keys())
    })




@app.route('/api/conversations', methods=['GET', 'POST'])
def conversations():
    print("Hit /api/conversations route with method", request.method)

    convos = load_conversations()

    if request.method == 'POST':
        new_id = str(uuid.uuid4())
        convos[new_id] = {
            "title": "Untitled",
            "messages": []
        }
        save_conversations(convos)
        return jsonify({"id": new_id})

    return jsonify(convos)

@app.route('/api/conversations/<conversation_id>', methods=['DELETE'])
def delete_conversation(conversation_id):
    convos = load_conversations()
    if conversation_id in convos:
        del convos[conversation_id]
        save_conversations(convos)
        return jsonify({"status": "deleted"})
    return jsonify({"error": "Conversation not found"}), 404

@app.route('/api/update-title', methods=['POST'])
def update_title():
    data = request.get_json()
    conversation_id = data.get("conversationId")
    messages = data.get("messages", [])
    convos = load_conversations()

    if not conversation_id or len(messages) < 2:
        return jsonify({"status": "ignored"})

    if convos[conversation_id]["title"] != "Untitled":
        return jsonify({"status": "title_already_set", "title": convos[conversation_id]["title"]})

    prompt_lines = [
        "Generate a concise and professional title based on the following conversation.",
        "Do NOT use quotation marks or emojis.",
        "Examples: 'Mathematics Problem', 'AI Ethics', 'Software Development Help'",
        "\n---\nConversation:\n"
    ]
    for msg in messages[:5]:
        role = "User" if msg["role"] == "user" else "Assistant"
        prompt_lines.append(f"{role}: {msg['content']}")
    prompt = "\n".join(prompt_lines)

    try:
        response = requests.post(OLLAMA_URL, json={
            "model": MODEL_NAME,
            "messages": [
                {"role": "system", "content": "You are a helpful assistant."},
                {"role": "user", "content": prompt}
            ],
            "stream": False
        })
        if response.ok:
            title = response.json()["message"]["content"].strip()
            if title:
                convos[conversation_id]["title"] = title
                save_conversations(convos)
                return jsonify({"status": "title_updated", "title": title})
    except Exception as e:
        print(f"Error generating title: {e}")

    return jsonify({"status": "title_update_failed"})

@app.route('/api/edit-message', methods=['POST'])
def edit_message():
    data = request.get_json()
    conversation_id = data.get("conversationId")
    message_index = data.get("messageIndex")
    new_content = data.get("newContent")
    has_image = data.get("hasImage", False)  # Check if the message has an image

    if not all([conversation_id, message_index is not None, new_content]):
        return jsonify({"error": "Missing data"}), 400

    convos = load_conversations()
    if conversation_id not in convos:
        return jsonify({"error": "Conversation not found"}), 404

    messages = convos[conversation_id]["messages"]

    if message_index >= len(messages) or messages[message_index]["role"] != "user":
        return jsonify({"error": "Invalid message index or not a user message"}), 400

    if message_index + 1 >= len(messages) or messages[message_index + 1]["role"] != "assistant":
        return jsonify({"error": "No assistant reply to update"}), 400

    # Truncate everything after this user + assistant pair
    del messages[message_index + 2:]

    # Update user message
    messages[message_index]["content"] = new_content

    # Regenerate assistant reply
    try:
        # Use llava model if the message has an image, otherwise use regular model
        model_to_use = "llava" if has_image else MODEL_NAME
        
        # Prepare messages for the API call
        messages_for_api = messages[:message_index + 1]
        
        # If it's an image message, we need to preserve the image data
        if has_image and messages[message_index].get("image"):
            # For image messages, we need to reconstruct the proper format
            # The frontend stores the image as a data URL, but we need base64 for Ollama
            image_data_url = messages[message_index]["image"]
            
            # Extract base64 data from data URL
            if image_data_url.startswith('data:image'):
                # Remove the data URL prefix to get just the base64 data
                base64_data = image_data_url.split(',')[1]
                
                # Create the proper message format for llava
                messages_for_api = [{
                    "role": "user",
                    "content": new_content,
                    "images": [base64_data]
                }]
            else:
                # If it's already base64 data (from old format), use it directly
                base64_data = image_data_url
                messages_for_api = [{
                    "role": "user",
                    "content": new_content,
                    "images": [base64_data]
                }]
        else:
            # Fallback to regular text if image data is corrupted
            has_image = False
            model_to_use = MODEL_NAME
        
        response = requests.post(OLLAMA_URL, json={
            "model": model_to_use,
            "messages": messages_for_api,
            "stream": False
        })
        ai_reply = response.json()["message"]["content"] if response.ok else "❌ Error from LLM."
    except Exception as e:
        print(f"[ERROR] Error in edit_message: {e}")
        ai_reply = "⚠️ Unable to reach backend."

    # Replace AI message or append if missing
    if message_index + 1 < len(messages):
        messages[message_index + 1]["content"] = ai_reply
    else:
        messages.append({"role": "assistant", "content": ai_reply})

    save_conversations(convos)
    return jsonify({
        "content": ai_reply,
        "conversation": convos[conversation_id]
    })


@app.route('/api/chat-with-image', methods=['POST'])
def chat_with_image():
    print(f"[DEBUG] Image chat endpoint called")
    print(f"[DEBUG] Request files: {list(request.files.keys())}")
    print(f"[DEBUG] Request form: {list(request.form.keys())}")
    
    if 'image' not in request.files:
        print(f"[ERROR] No image in request files")
        return jsonify({'error': 'No image provided'}), 400

    image_file = request.files['image']
    prompt = request.form.get('prompt', '')  # Optional user message
    conversation_id = request.form.get('conversationId')  # You need to send this from frontend
    
    print(f"[DEBUG] Image file: {image_file.filename}")
    print(f"[DEBUG] Prompt: {prompt}")
    print(f"[DEBUG] Conversation ID: {conversation_id}")

    if not conversation_id:
        print(f"[ERROR] No conversation ID provided")
        return jsonify({"error": "No conversation ID provided"}), 400
    
    # Validate conversation ID format
    try:
        conversation_uuid = uuid.UUID(str(conversation_id))
        conversation_id = str(conversation_uuid)  # Ensure it's a string
        print(f"[DEBUG] Validated conversation ID: {conversation_id}")
    except ValueError:
        print(f"[ERROR] Invalid conversation ID format: {conversation_id}")
        return jsonify({"error": "Invalid conversation ID format"}), 400

    # Read image and encode as base64
    image_data = image_file.read()
    encoded_image = base64.b64encode(image_data).decode("utf-8")

    # Prepare messages for Ollama
    messages = [{
        "role": "user",
        "content": prompt,
        "images": [encoded_image]
    }]

    # Send to Ollama using synchronous requests
    try:
        response = requests.post(OLLAMA_URL, json={
            "model": "llava",  # Use a vision-capable model
            "messages": messages,
            "stream": False
        }, timeout=60)
        
        if response.status_code == 200:
            result = response.json()
            print("LLM Response:", result)
            reply = result.get("message", {}).get("content", "")

            # Save conversation update here without touching the rest:
            convos = load_conversations()
            print(f"[DEBUG] Available conversations: {list(convos.keys())}")
            print(f"[DEBUG] Looking for conversation: {conversation_id}")
            print(f"[DEBUG] Conversation ID type: {type(conversation_id)}")
            print(f"[DEBUG] Available conversation IDs types: {[type(k) for k in convos.keys()]}")
            
            # Try to find the conversation with different key types
            found_conversation = None
            for key in convos.keys():
                if str(key) == str(conversation_id):
                    found_conversation = key
                    break
            
            if found_conversation is None:
                print(f"[ERROR] Conversation {conversation_id} not found")
                return jsonify({"error": "Conversation not found"}), 404
            
            conversation_id = found_conversation  # Use the actual key from the conversations dict

            # Convert base64 to data URL format for frontend compatibility
            data_url = f"data:image/jpeg;base64,{encoded_image}"
            convos[conversation_id]["messages"].append({
                "role": "user",
                "content": prompt,
                "image": data_url
            })
            convos[conversation_id]["messages"].append({
                "role": "assistant",
                "content": reply
            })

            print(f"[DEBUG] Saving conversation with {len(convos[conversation_id]['messages'])} messages")
            save_conversations(convos)

            print(f"[DEBUG] Successfully returning response: {reply[:100]}...")
            print(f"[DEBUG] Response status: 200")
            return jsonify({"content": reply})
        else:
            error = response.text
            print(f"[ERROR] Ollama API error: {error}")
            print(f"[ERROR] Response status: {response.status_code}")
            return jsonify({"error": "Ollama API error", "details": error}), 500
    except Exception as e:
        print(f"[ERROR] Exception in image chat: {e}")
        print(f"[ERROR] Exception type: {type(e)}")
        import traceback
        print(f"[ERROR] Traceback: {traceback.format_exc()}")
        return jsonify({"error": "Error calling Ollama", "exception": str(e)}), 500















# ========== HELPERS ==========


# ========== START SERVER ==========
# Start the background thread for preloading files
file_cache_thread = threading.Thread(target=background_file_cache_refresher, args=(3600,), daemon=True)
file_cache_thread.start()

if __name__ == '__main__':
    app.run(debug=True, port=5000)
